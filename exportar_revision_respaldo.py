"""
Respaldo del Excel de revisión con openpyxl (ya incluido en la app).

Se usa solo si XlsxWriter no está instalado en el servidor, para que la app
nunca se caiga por una librería faltante. Diferencia: aquí las fórmulas no
guardan su resultado, así que en vistas previas (Mac, correo, Drive) algunas
celdas pueden verse vacías; en Excel se calculan al abrir.
"""

from __future__ import annotations

import datetime as _dt
import io

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import validacion_referencias as v

FUENTES = ("historico", "nl", "cdmx", "ia")

AZUL = "1F3864"
AZUL2 = "2E5496"
GRIS = "F2F2F2"
ROJO_F, ROJO_T = "F8C9C9", "712121"
VERDE_F, VERDE_T = "CCEBD2", "14532D"
AMBAR_F, AMBAR_T = "FFF0BB", "705000"
GRIS_F, GRIS_T = "E4E7EB", "475467"
INPUT_T = "0000FF"

_fino = Side(style="thin", color="BFBFBF")
BORDE = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
MONEDA = '$#,##0.00;($#,##0.00);"—"'
PCT = '+0.0%;-0.0%;0.0%'


def _encabezado(ws, fila, textos, relleno=AZUL2):
    for col, texto in enumerate(textos, start=1):
        c = ws.cell(row=fila, column=col, value=texto)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=relleno)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE


def _titulo(ws, texto, subtitulo, ancho):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)
    c = ws.cell(row=1, column=1, value=texto)
    c.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=AZUL)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho)
    c = ws.cell(row=2, column=1, value=subtitulo)
    c.font = Font(name="Arial", size=9, color="595959")


def _celda(ws, fila, col, valor, fmt=None, azul=False, negrita=False, wrap=False, centro=False):
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = Font(name="Arial", size=10, bold=negrita, color=INPUT_T if azul else "000000")
    c.border = BORDE
    c.alignment = Alignment(vertical="top", wrap_text=wrap,
                            horizontal="center" if centro else None)
    if fmt:
        c.number_format = fmt
    return c


def _semaforo_cf(ws, rango):
    for texto, fondo, letra in (
        ("ALTO", ROJO_F, ROJO_T),
        ("BAJO", VERDE_F, VERDE_T),
        ("EN MERCADO", AMBAR_F, AMBAR_T),
        ("MIXTO", GRIS_F, GRIS_T),
        ("NO CONCLUYENTE", GRIS_F, GRIS_T),
        ("SIN DATOS SUFICIENTES", GRIS_F, GRIS_T),
    ):
        ws.conditional_formatting.add(
            rango,
            CellIsRule(operator="equal", formula=[f'"{texto}"'],
                       fill=PatternFill("solid", fgColor=fondo),
                       font=Font(color=letra, bold=True)),
        )


def generar_excel(filas: list[dict], proveedor: str = "", proyecto: str = "", **_ignorados) -> bytes:  # respaldo openpyxl
    """filas: los dicts de la app (con '_evaluaciones' y '_final')."""
    wb = Workbook()

    # ------------------------------------------------------------------
    # Hoja 2 primero (Resumen la referencia): Por fuente
    # ------------------------------------------------------------------
    ws_f = wb.active
    ws_f.title = "Por fuente"
    cab_f = ["#", "Concepto", "Unidad", "P.U. cotizado"]
    for clave in FUENTES:
        nombre = v.NOMBRE_FUENTE[clave]
        cab_f += [f"{nombre}\nprecio ref.", f"{nombre}\nestado",
                  f"{nombre}\n% vs cotizado", f"{nombre}\nsemáforo"]
    _titulo(ws_f, "Comparación independiente por fuente",
            "Cada fuente se compara por separado (sin promediar). Solo las referencias VALIDADAS "
            "tienen semáforo. Semáforo: ±5 % = EN MERCADO.", len(cab_f))
    _encabezado(ws_f, 4, cab_f)
    ws_f.row_dimensions[4].height = 42
    fila_f = {}
    for i, f in enumerate(filas):
        r = 5 + i
        fila_f[i] = r
        _celda(ws_f, r, 1, f.get("Partida"), centro=True)
        _celda(ws_f, r, 2, f.get("Concepto"), wrap=True)
        _celda(ws_f, r, 3, f.get("Unidad"), centro=True)
        _celda(ws_f, r, 4, f.get("Precio cotizado"), MONEDA, azul=True)
        evs = f.get("_evaluaciones") or {}
        for j, clave in enumerate(FUENTES):
            base = 5 + j * 4
            ev = evs.get(clave) or {}
            ref_col = get_column_letter(base)
            est_col = get_column_letter(base + 1)
            precio_ref = ev.get("precio_referencia")
            _celda(ws_f, r, base, precio_ref, MONEDA, azul=True)
            _celda(ws_f, r, base + 1, ev.get("estado", v.SIN_DATO), centro=True)
            # Las rechazadas y las sin dato no tienen % ni semáforo.
            _celda(ws_f, r, base + 2,
                   f'=IF(OR({ref_col}{r}="",{est_col}{r}="RECHAZADA",{est_col}{r}="SIN DATO"),"",D{r}/{ref_col}{r}-1)',
                   PCT)
            _celda(ws_f, r, base + 3,
                   f'=IF({est_col}{r}<>"VALIDADA","",IF(D{r}>{ref_col}{r}*1.05,"ALTO",'
                   f'IF(D{r}<{ref_col}{r}*0.95,"BAJO","EN MERCADO")))',
                   centro=True)
            _semaforo_cf(ws_f, f"{get_column_letter(base + 3)}{r}")
    ultima_f = 4 + len(filas)
    for j in range(len(FUENTES)):
        col_est = get_column_letter(6 + j * 4)
        rng = f"{col_est}5:{col_est}{max(ultima_f, 5)}"
        for texto, color in (("VALIDADA", VERDE_T), ("RECHAZADA", ROJO_T),
                             ("NO CONCLUYENTE", AMBAR_T), ("POR CONFIRMAR", GRIS_T)):
            ws_f.conditional_formatting.add(
                rng, CellIsRule(operator="equal", formula=[f'"{texto}"'],
                                font=Font(color=color, bold=True)))
    anchos_f = {1: 5, 2: 44, 3: 8, 4: 13}
    for j in range(len(FUENTES)):
        anchos_f.update({5 + j * 4: 13, 6 + j * 4: 15, 7 + j * 4: 11, 8 + j * 4: 12})
    for col, ancho in anchos_f.items():
        ws_f.column_dimensions[get_column_letter(col)].width = ancho
    ws_f.freeze_panes = "E5"
    ws_f.sheet_view.showGridLines = False

    # ------------------------------------------------------------------
    # Hoja 1: Resumen
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Resumen", 0)
    cab = ["#", "Concepto", "Unidad", "Cantidad", "P.U. cotizado", "Importe cotizado",
           "Semáforo final", "Fuentes validadas", "Referencia para negociar",
           "P.U. de negociación", "% vs negociación", "Ahorro potencial", "Nota"]
    _titulo(ws, "Revisión de cotización CAPEX",
            f"Proveedor: {proveedor or '—'} · Proyecto: {proyecto or '—'} · "
            f"Generado {_dt.date.today():%d/%m/%Y}. Azul = dato capturado; negro = fórmula.",
            len(cab))
    _encabezado(ws, 4, cab)
    ws.row_dimensions[4].height = 32
    for i, f in enumerate(filas):
        r = 5 + i
        fin = f.get("_final") or {}
        _celda(ws, r, 1, f.get("Partida"), centro=True)
        _celda(ws, r, 2, f.get("Concepto"), wrap=True)
        _celda(ws, r, 3, f.get("Unidad"), centro=True)
        _celda(ws, r, 4, f.get("Cantidad"), "#,##0.00", azul=True)
        _celda(ws, r, 5, f.get("Precio cotizado"), MONEDA, azul=True)
        _celda(ws, r, 6, f"=D{r}*E{r}", MONEDA)
        _celda(ws, r, 7, fin.get("semaforo") or v.SIN_VALIDADA, centro=True, negrita=True)
        _celda(ws, r, 8, fin.get("fuentes_validadas") or "—", wrap=True)
        _celda(ws, r, 9, fin.get("referencia_negociacion") or "—")
        # El P.U. de negociación se enlaza a la hoja "Por fuente" (misma
        # referencia que la app eligió), para que cambie si se corrige ahí.
        ref = fin.get("referencia_negociacion")
        clave_ref = next((k for k, n in v.NOMBRE_FUENTE.items() if n == ref), None)
        if clave_ref:
            col_ref = get_column_letter(5 + FUENTES.index(clave_ref) * 4)
            _celda(ws, r, 10, f"='Por fuente'!{col_ref}{fila_f[i]}", MONEDA)
        else:
            _celda(ws, r, 10, None, MONEDA)
        _celda(ws, r, 11, f'=IF(J{r}="","",E{r}/J{r}-1)', PCT)
        _celda(ws, r, 12, f'=IF(J{r}="",0,MAX(0,(E{r}-J{r})*D{r}))', MONEDA)
        _celda(ws, r, 13, fin.get("detalle") or "", wrap=True)
    ultima = 4 + len(filas)
    rt = ultima + 1
    _celda(ws, rt, 2, "TOTAL", negrita=True)
    _celda(ws, rt, 6, f"=SUM(F5:F{ultima})", MONEDA, negrita=True)
    _celda(ws, rt, 12, f"=SUM(L5:L{ultima})", MONEDA, negrita=True)
    _celda(ws, rt + 1, 2, "Ahorro potencial sobre el importe", negrita=True)
    _celda(ws, rt + 1, 12, f'=IF(F{rt}=0,"",L{rt}/F{rt})', "0.0%", negrita=True)
    _semaforo_cf(ws, f"G5:G{max(ultima, 5)}")
    for col, ancho in {1: 5, 2: 46, 3: 8, 4: 10, 5: 13, 6: 15, 7: 23, 8: 22,
                       9: 19, 10: 15, 11: 12, 12: 15, 13: 40}.items():
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = "C5"
    ws.sheet_view.showGridLines = False

    # ------------------------------------------------------------------
    # Hoja 3: Evidencia
    # ------------------------------------------------------------------
    ws_e = wb.create_sheet("Evidencia")
    cab_e = ["#", "Concepto cotizado", "Unidad", "Fuente", "Estado", "Motivo / verificaciones",
             "Concepto encontrado en la fuente", "Confianza texto", "P.U. referencia",
             "Revisión IA", "Evidencia web (frase o URL)"]
    _titulo(ws_e, "Evidencia de cada referencia",
            "Por qué cada referencia quedó VALIDADA, NO CONCLUYENTE, POR CONFIRMAR, RECHAZADA o SIN DATO.",
            len(cab_e))
    _encabezado(ws_e, 4, cab_e)
    r = 5
    for f in filas:
        for clave in FUENTES:
            ev = (f.get("_evaluaciones") or {}).get(clave) or {}
            _celda(ws_e, r, 1, f.get("Partida"), centro=True)
            _celda(ws_e, r, 2, f.get("Concepto"), wrap=True)
            _celda(ws_e, r, 3, f.get("Unidad"), centro=True)
            _celda(ws_e, r, 4, v.NOMBRE_FUENTE[clave])
            _celda(ws_e, r, 5, ev.get("estado", v.SIN_DATO), centro=True, negrita=True)
            _celda(ws_e, r, 6, ev.get("motivo"), wrap=True)
            _celda(ws_e, r, 7, ev.get("descripcion"), wrap=True)
            _celda(ws_e, r, 8, ev.get("confianza"), centro=True)
            _celda(ws_e, r, 9, ev.get("precio_referencia"), MONEDA)
            _celda(ws_e, r, 10, ev.get("revision_ia"), wrap=True)
            _celda(ws_e, r, 11, ev.get("evidencia"), wrap=True)
            r += 1
    rng_e = f"E5:E{max(r - 1, 5)}"
    for texto, color in (("VALIDADA", VERDE_T), ("RECHAZADA", ROJO_T),
                         ("NO CONCLUYENTE", AMBAR_T), ("POR CONFIRMAR", GRIS_T)):
        ws_e.conditional_formatting.add(
            rng_e, CellIsRule(operator="equal", formula=[f'"{texto}"'], font=Font(color=color, bold=True)))
    for col, ancho in {1: 5, 2: 38, 3: 8, 4: 16, 5: 16, 6: 44, 7: 50, 8: 11, 9: 14,
                       10: 44, 11: 50}.items():
        ws_e.column_dimensions[get_column_letter(col)].width = ancho
    ws_e.freeze_panes = "E5"
    ws_e.sheet_view.showGridLines = False

    # ------------------------------------------------------------------
    # Hoja 4: Metodología
    # ------------------------------------------------------------------
    ws_m = wb.create_sheet("Metodología")
    ws_m.column_dimensions["A"].width = 26
    ws_m.column_dimensions["B"].width = 100
    _titulo(ws_m, "Metodología de validación", "Cómo se decide qué referencias cuentan.", 2)
    reglas = [
        ("RECHAZA (IA)", "La referencia se excluye del precio de negociación, de las diferencias y del semáforo. Queda solo como evidencia."),
        ("NO_SEGURO (IA)", "Se conserva como orientativa (NO CONCLUYENTE): se muestra su precio y su %, pero no decide el semáforo."),
        ("CONFIRMA (IA)", "Además se verifica unidad (misma unidad), alcance (suministro vs. instalación), antigüedad del dato (máx. "
                          f"{v.ANTIGUEDAD_MAXIMA_ANIOS} años) y escala del precio (dentro de {v.FACTOR_ESCALA:.0f}× del cotizado)."),
        ("Coincidencia débil", "Si la coincidencia de texto es BAJA y la IA no la confirmó: POR CONFIRMAR (orientativa)."),
        ("Precio web", "Debe aparecer en la misma frase que el concepto y la unidad. Un fragmento web sin validar es orientativo; "
                       "solo un precio de Gemini con fuente y frase verificadas puede quedar VALIDADO."),
        ("Sin promedios", "Cada fuente se compara por separado. El P.U. de negociación sale de UNA referencia validada, con prioridad: "
                          "Histórico Ragasa > Nuevo León > CDMX > IA internet."),
        ("Semáforo final", "Todas las validadas coinciden = ese semáforo. Si discrepan = MIXTO. Sin validadas pero con orientativas = "
                           "NO CONCLUYENTE. Sin nada = SIN DATOS SUFICIENTES. EN MERCADO = ±5 % de la referencia."),
        ("Nuevo León", "Mediana del tabulador homologado de licitaciones de NL, actualizada a hoy con INPC (INEGI)."),
        ("CDMX", "Tabulador General de Precios Unitarios del Gobierno de la CDMX."),
        ("Histórico Ragasa", "Cotizaciones guardadas previamente en el histórico interno (Google Sheets)."),
    ]
    _encabezado(ws_m, 4, ["Regla", "Aplicación"])
    for i, (regla, texto) in enumerate(reglas):
        _celda(ws_m, 5 + i, 1, regla, negrita=True)
        _celda(ws_m, 5 + i, 2, texto, wrap=True)
        ws_m.row_dimensions[5 + i].height = 32
    ws_m.sheet_view.showGridLines = False

    for hoja in wb.worksheets:
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 0
        hoja.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
